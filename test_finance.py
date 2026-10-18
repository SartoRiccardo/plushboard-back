"""Finance module + routes check (multi-leg double-entry).
Run from backend/: .venv/bin/python test_finance.py"""
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient

from app import db
from app.auth import serializer

db.DB_PATH = Path(tempfile.mkdtemp()) / "t.db"
db.init_db()

from app.main import app  # noqa: E402

COOKIE = {"plush_session": serializer.dumps({"email": "rikki.sarto@gmail.com"})}

with TestClient(app) as c:
    # accounts: Bank (asset, grows with Dare), Sales + Fees (revenue/expense)
    ids = {}
    for name, dare in (("Bank", True), ("Sales", False), ("Fees", True)):
        r = c.post("/api/finance/accounts", json={"name": name, "dare_increases": dare},
                   cookies=COOKIE)
        assert r.status_code == 200, r.text
        ids[name] = r.json()["id"]
    bal = {a["name"]: a["balance"] for a in
           c.get("/api/finance/accounts", cookies=COOKIE).json()}
    assert bal == {"Bank": 0, "Sales": 0, "Fees": 0}

    # balanced 2-leg EUR movement
    r = c.post("/api/finance/movements", json={
        "date": "2026-10-01", "description": "sale", "currency": "EUR",
        "legs": [{"account_id": ids["Bank"], "dare": 100},
                 {"account_id": ids["Sales"], "avere": 100}]}, cookies=COOKIE)
    assert r.status_code == 200, r.text
    assert r.json()["rate"] is None and r.json()["eur_amount"] == 200

    # 3-leg USD movement with manual rate: 59.97 dare vs 41.97+18 avere
    r = c.post("/api/finance/movements", json={
        "date": "2026-10-02", "description": "usd sale", "currency": "USD",
        "manual_rate": 0.9,
        "legs": [{"account_id": ids["Bank"], "dare": 59.97},
                 {"account_id": ids["Sales"], "avere": 41.97},
                 {"account_id": ids["Fees"], "avere": 18.0}]}, cookies=COOKIE)
    assert r.status_code == 200, r.text
    assert r.json()["rate"] == 0.9

    # unbalanced 3-leg → 400
    r = c.post("/api/finance/movements", json={
        "date": "2026-10-01", "description": "bad", "currency": "EUR",
        "legs": [{"account_id": ids["Bank"], "dare": 100},
                 {"account_id": ids["Sales"], "avere": 40},
                 {"account_id": ids["Fees"], "avere": 40}]}, cookies=COOKIE)
    assert r.status_code == 400 and "unbalanced" in r.text
    # bad currency / negative leg / unknown account → 400
    # (USD without manual_rate is valid: it fetches the live ECB rate)
    base = {"date": "2026-10-01", "description": "x"}
    for body in [
        {**base, "currency": "GBP",
         "legs": [{"account_id": ids["Bank"], "dare": 1}, {"account_id": ids["Fees"], "avere": 1}]},
        {**base, "currency": "EUR",
         "legs": [{"account_id": ids["Bank"], "dare": -1}, {"account_id": ids["Fees"], "avere": 1}]},
        {**base, "currency": "EUR",
         "legs": [{"account_id": 9999, "dare": 1}, {"account_id": ids["Fees"], "avere": 1}]},
        {**base, "currency": "EUR", "legs": [{"account_id": ids["Bank"], "dare": 1}]},
    ]:
        r = c.post("/api/finance/movements", json=body, cookies=COOKIE)
        assert r.status_code == 400, (body, r.text)

    # balances: Bank dare = 100 + 53.97; Sales avere = 100 + 37.77;
    # Fees has only avere legs → negative (expense account credited)
    bal = {a["name"]: round(a["balance"], 2) for a in
           c.get("/api/finance/accounts", cookies=COOKIE).json()}
    assert bal["Bank"] == 153.97 and bal["Sales"] == 137.77 and bal["Fees"] == -16.2

    # movements page: aggregation + totals
    r = c.get("/api/finance/movements", cookies=COOKIE)
    items = r.json()["items"]
    assert r.json()["total"] == 2 and r.json()["pages"] == 1
    usd = next(i for i in items if i["description"] == "usd sale")
    assert usd["dare_accounts"] == ["Bank"]
    assert sorted(usd["avere_accounts"]) == ["Fees", "Sales"]
    assert round(usd["eur_amount"], 2) == 107.94
    ov = c.get("/api/finance/overview", cookies=COOKIE).json()
    assert round(ov["totals"]["dare_eur"], 2) == round(ov["totals"]["avere_eur"], 2)
    assert ov["totals"]["movements"] == 2

    # importer grouping regression: the owner splits one Ko-Fi sale across
    # rows with suffixed descriptions — legs must accumulate into ONE
    # movement that balances only when summed (grouping by exact description
    # used to import these as 4 unbalanced movements)
    from import_xlsx import _groups, _unbalanced, SPECIAL_CASES, import_file

    def leg(rn, dare, avere, eur, descr):
        return (rn, ("2026-05-23", "C", dare, avere, "EUR",
                     eur if dare else 0, 0 if dare else eur, descr))

    kofi = [  # dare legs (goods + fees) vs avere legs (sale + inventory), 6 rows
        leg(2, 65.87, 0, 56.74, "Ko-Fi S-W5E720147D"),
        leg(3, 2.90, 0, 2.50, "Ko-Fi S-W5E720147D (Ko-Fi fee)"),
        leg(4, 4.23, 0, 3.64, "Ko-Fi S-W5E720147D (PP fee)"),
        leg(5, 12.08, 0, 10.41, "Ko-Fi S-W5E720147D"),
        leg(6, 0, 12.08, 10.41, "Ko-Fi S-W5E720147D"),
        leg(7, 0, 73.00, 62.88, "Ko-Fi S-W5E720147D"),
    ]
    groups, leftover = _groups(kofi)
    assert len(groups) == 1 and not leftover
    assert [rn for rn, _ in groups[0]] == [2, 3, 4, 5, 6, 7]
    dare = sum(r[2] for _, r in groups[0])
    avere = sum(r[3] for _, r in groups[0])
    assert round(dare, 2) == round(avere, 2) == 85.08  # balanced only when summed
    assert groups[0][0][1][7] == "Ko-Fi S-W5E720147D"  # first leg's descr, verbatim

    # end-to-end: SPECIAL_CASES force-group exact rows into one movement;
    # legs that never balance are printed as leftovers and NOT imported, and
    # the account (Bilancio) gate aborts — nothing is committed
    import openpyxl

    def write_xlsx(rows, conti):
        w = openpyxl.Workbook()
        ws = w.active
        ws.title = "BK@Conti"
        ws.append(("Nome conto", "Descrizione", "Bilancio", "Cresce con Avere"))
        for row in conti:
            ws.append(row)
        t = w.create_sheet("BK@Transazioni")
        t.append(("Data", "Conto", "Dare", "Avere", "Valuta",
                  "Dare (EUR)", "Avere (EUR)", "Descrizione"))
        for row in rows:
            t.append(row)
        w.create_sheet("BK@_conversion_to_eur")
        path = Path(tempfile.mkdtemp()) / "ledger.xlsx"
        w.save(path)
        return str(path)

    spec = {"rows": [2, 3, 4], "description": "vendita x"}
    head = [("2026-06-01", "A", 85.65, 0, "EUR", 85.65, 0, "vendita x (fee)"),
            ("2026-06-01", "B", 0, 3.72, "EUR", 0, 3.72, "vendita x (pp)"),
            ("2026-06-01", "C", 0, 81.93, "EUR", 0, 81.93, "vendita x")]
    SPECIAL_CASES.append(spec)
    try:
        xlsx = write_xlsx(
            head + [("2026-06-01", "A", 12.0, 0, "EUR", 12.0, 0, "solo dare")],
            [("A", "", 97.65, False), ("B", "", 3.72, True), ("C", "", 81.93, True)])
        try:
            import_file(xlsx)
            raise AssertionError("leftover legs must fail the Bilancio gate")
        except SystemExit as e:
            assert "account(s) off" in str(e)
        with db.db() as dbc:
            n = dbc.execute("SELECT COUNT(*) AS n FROM movements").fetchone()["n"]
            assert n == 2  # the aborted import rolled back (API movements only)

        xlsx = write_xlsx(  # add the missing counterpart: no leftover, clean run
            head + [("2026-06-02", "A", 12.0, 0, "EUR", 12.0, 0, "pagamento 1"),
                    ("2026-06-02", "B", 0, 12.0, "EUR", 0, 12.0, "pagamento 2")],
            [("A", "", 97.65, False), ("B", "", 15.72, True), ("C", "", 81.93, True)])
        assert [r["account"] for r in import_file(xlsx)] == ["A", "B", "C"]
        with db.db() as dbc:
            descs = [r["description"] for r in dbc.execute(
                "SELECT description FROM movements WHERE source='import'"
                " ORDER BY id")]
            assert descs == ["vendita x", "pagamento 1"]  # spec descr, then first leg's
            assert _unbalanced(dbc) == []  # every imported movement balances
    finally:
        SPECIAL_CASES.clear()

    # delete movement; account with legs refuses deletion
    assert c.delete(f"/api/finance/movements/{items[0]['id']}",
                    cookies=COOKIE).status_code == 200
    r = c.delete(f"/api/finance/accounts/{ids['Bank']}", cookies=COOKIE)
    assert r.status_code == 400 and "movements" in r.text

    # products CRUD
    r = c.post("/api/finance/products",
               json={"name": "Mini Nareko Plush", "unit_price": 48.0}, cookies=COOKIE)
    pid = r.json()["id"]
    r = c.post("/api/finance/products", json={"name": "", "unit_price": 1}, cookies=COOKIE)
    assert r.status_code == 400
    assert c.put(f"/api/finance/products/{pid}", json={"unit_price": 41.17},
                 cookies=COOKIE).status_code == 200
    assert c.get("/api/finance/products", cookies=COOKIE).json()[0]["unit_price"] == 41.17
    assert c.delete(f"/api/finance/products/{pid}", cookies=COOKIE).status_code == 200
    assert c.delete(f"/api/finance/products/{pid}", cookies=COOKIE).status_code == 404

print("finance ok")
